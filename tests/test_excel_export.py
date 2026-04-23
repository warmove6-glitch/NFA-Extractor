"""
Testes para o módulo excel_export.py.
Valida a geração das 4 abas do Excel e o comportamento de borda.
"""

import sys
import tempfile
from pathlib import Path

import pytest
import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from extractor import NFA, Parte, Produto
from excel_export import (
    exportar_excel,
    _cel_style,
    _cor_linha,
    _aba_notas,
    _aba_itens,
    _aba_destinatarios,
    _aba_mensal,
)
from constants import CORES


# ─── Fixtures locais ──────────────────────────────────────────────────────────

@pytest.fixture
def notas_exportacao() -> list[NFA]:
    """Conjunto de notas para testar exportação completa."""
    return [
        NFA(
            numero='100001',
            emissao='15/03/2025',
            natureza='VENDA DE GADO BOVINO',
            local_emissao='AGENCIA FAZENDARIA DE TROMBAS',
            destinatario=Parte(nome='FRIGORIFICO ALPHA', cpf_cnpj='11.111.111/0001-11', municipio='GOIANIA'),
            transportador=Parte(nome='TRANSPORTADORA XPTO'),
            produtos=[
                Produto(codigo='1070', descricao='NELORE MACHO ATE 12M', quantidade=10.0,
                        vlr_unitario=2500.0, vlr_total=25000.0),
                Produto(codigo='629', descricao='NELORE FEMEA ACIMA 36M', quantidade=5.0,
                        vlr_unitario=3300.0, vlr_total=16500.0),
            ],
        ),
        NFA(
            numero='100002',
            emissao='20/03/2025',
            natureza='REMESSA DE BEZERROS',
            local_emissao='NOTA EMITIDA PELO PRÓPRIO CONTRIBUINTE',
            destinatario=Parte(nome='FAZENDA BOA VISTA', cpf_cnpj='22.222.222/0001-22', municipio='FORMOSO'),
            produtos=[Produto(codigo='1071', descricao='NELORE MACHO 13-24M', quantidade=8.0,
                              vlr_unitario=3000.0, vlr_total=24000.0)],
        ),
        NFA(
            numero='100003',
            emissao='10/04/2025',
            natureza='VENDA DE GADO BOVINO',
            local_emissao='AGENCIA FAZENDARIA DE FORMOSO',
            destinatario=Parte(nome='FRIGORIFICO ALPHA', cpf_cnpj='11.111.111/0001-11', municipio='GOIANIA'),
            produtos=[Produto(codigo='1077', descricao='NELORE FEMEA 13-24M', quantidade=12.0,
                              vlr_unitario=2800.0, vlr_total=33600.0)],
        ),
    ]


# ─── Testes: _cor_linha() ─────────────────────────────────────────────────────

class TestCorLinha:
    """Valida a alternância de cores por linha."""

    def test_linha_par_retorna_alt(self):
        assert _cor_linha(2) == CORES['ALT']
        assert _cor_linha(4) == CORES['ALT']

    def test_linha_impar_retorna_white(self):
        assert _cor_linha(1) == CORES['WHITE']
        assert _cor_linha(3) == CORES['WHITE']

    def test_linha_zero_retorna_alt(self):
        # 0 é par
        assert _cor_linha(0) == CORES['ALT']


# ─── Testes: _cel_style() ────────────────────────────────────────────────────

class TestCelStyle:
    """Valida aplicação de estilo a células."""

    def test_valor_escrito_corretamente(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 'Teste')
        assert cell.value == 'Teste'

    def test_formato_numerico_aplicado(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 1234.56, num_fmt='#,##0.00')
        assert cell.number_format == '#,##0.00'

    def test_bold_aplicado(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 'Header', bold=True)
        assert cell.font.bold is True

    def test_sem_bold_por_padrao(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 'Normal')
        assert cell.font.bold is False

    def test_cor_fonte_aplicada(self):
        """openpyxl armazena RGB como ARGB (8 chars com prefixo 00)."""
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 'X', fg=CORES['CYAN'])
        # openpyxl normaliza para ARGB (ex: "00D4FF" → "0000D4FF")
        assert cell.font.color.rgb.endswith(CORES['CYAN'])

    def test_alinhamento_direita(self):
        wb = openpyxl.Workbook()
        ws = wb.active
        cell = _cel_style(ws, 1, 1, 100.0, align='right')
        assert cell.alignment.horizontal == 'right'


# ─── Testes: exportar_excel() — arquivo gerado ───────────────────────────────

class TestExportarExcel:
    """Testa geração completa do arquivo Excel."""

    def test_arquivo_criado(self, notas_exportacao):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        assert Path(saida).exists()
        assert Path(saida).stat().st_size > 0

    def test_quatro_abas_geradas(self, notas_exportacao):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        assert len(wb.sheetnames) == 4

    def test_nomes_das_abas(self, notas_exportacao):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        assert 'Notas Fiscais' in wb.sheetnames
        assert 'Itens Detalhados' in wb.sheetnames
        assert 'Por Destinatário' in wb.sheetnames
        assert 'Evolução Mensal' in wb.sheetnames

    def test_aba_notas_linhas_corretas(self, notas_exportacao):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Notas Fiscais']
        # 1 linha de cabeçalho + 3 notas
        linhas_dados = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(c for c in r)]
        assert len(linhas_dados) == 3

    def test_aba_itens_linhas_corretas(self, notas_exportacao):
        """3 notas com 2+1+1 produtos = 4 linhas de itens."""
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Itens Detalhados']
        linhas_dados = [r for r in ws.iter_rows(min_row=2, values_only=True) if any(c for c in r)]
        assert len(linhas_dados) == 4  # 2 + 1 + 1

    def test_aba_dest_agrupa_por_comprador(self, notas_exportacao):
        """Dois registros têm o mesmo destinatário → deve agregar em 1 linha."""
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Por Destinatário']
        # 2 destinatários distintos + 1 linha de TOTAL
        nomes = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value]
        assert 'FRIGORIFICO ALPHA' in nomes
        assert 'FAZENDA BOA VISTA' in nomes

    def test_aba_mensal_meses_presentes(self, notas_exportacao):
        """Notas em Mar/25 e Abr/25 devem gerar 2 linhas mensais."""
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Evolução Mensal']
        meses = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)
                 if ws.cell(row=r, column=1).value]
        assert '03/2025' in meses
        assert '04/2025' in meses

    def test_numero_da_nota_na_aba_notas(self, notas_exportacao):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Notas Fiscais']
        numeros = [ws.cell(row=r, column=1).value for r in range(2, 5)]
        assert '100001' in numeros

    def test_valor_total_correto_na_aba_notas(self, notas_exportacao):
        """Nota 100001 tem valor total 25000 + 16500 = 41500."""
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas_exportacao, saida)
        wb = openpyxl.load_workbook(saida)
        ws = wb['Notas Fiscais']
        # Coluna 10 = Valor Total
        valores = [ws.cell(row=r, column=10).value for r in range(2, 5)]
        assert 41500.0 in valores


# ─── Testes: edge cases ───────────────────────────────────────────────────────

class TestExportarExcelEdgeCases:
    """Testa comportamentos de borda."""

    def test_lista_vazia_lanca_erro(self):
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        with pytest.raises(ValueError, match='Nenhuma nota'):
            exportar_excel([], saida)

    def test_nota_sem_produtos(self):
        """Nota sem itens não deve quebrar a exportação."""
        notas = [NFA(numero='999', emissao='01/01/2025', natureza='VENDA',
                     destinatario=Parte(nome='COMPRADOR', cpf_cnpj='000.000.000-00'))]
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas, saida)
        assert Path(saida).exists()

    def test_nota_sem_destinatario(self):
        """NFA com destinatário vazio não deve quebrar."""
        notas = [NFA(numero='888', emissao='01/01/2025', natureza='VENDA',
                     produtos=[Produto(quantidade=1.0, vlr_total=1000.0)])]
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas, saida)
        assert Path(saida).exists()

    def test_uma_nota_gera_arquivo_valido(self):
        notas = [NFA(
            numero='777',
            emissao='05/06/2025',
            natureza='VENDA',
            destinatario=Parte(nome='COMPRADOR A', cpf_cnpj='111.111.111-11'),
            produtos=[Produto(quantidade=3.0, vlr_total=9000.0)],
        )]
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as f:
            saida = f.name
        exportar_excel(notas, saida)
        wb = openpyxl.load_workbook(saida)
        assert len(wb.sheetnames) == 4
